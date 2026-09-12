import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import threading
from datetime import date
from pathlib import Path
from io import BytesIO

import requests
from flask import Flask, jsonify, request, send_from_directory, Response

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "navratri"
DB_PATH = BASE_DIR / "navratri_data.db"

DEFAULT_PIN = "1234"
ADMIN_PIN = os.environ.get("ADMIN_PIN", DEFAULT_PIN)

app = Flask(__name__, static_folder=None)
_lock = threading.Lock()

GENDER_FEE = {"bahin": 50, "bhai": 100}
MOBILE_RE = re.compile(r"^\d{10}$")
AADHAR_RE = re.compile(r"^\d{12}$")

DEFAULT_SETTINGS = {
    "title": "ગરબા પાસ – નવરાત્રી મહોત્સવ ૨૦૨૬",
    "subtitle": "શ્રી ખેલૈયા મંડળ, વાસદ",
    "address": "",
    "date_from": "11-10-2026",
    "date_to": "20-11-2026",
    "background": "",
}

VALID_ROLES = {"operator", "authority", "admin"}
VALID_STATUSES = {"pending", "approved", "rejected"}
SUPABASE_URL = os.environ.get("SUPABASE_URL", "").strip()
SUPABASE_KEY = os.environ.get("SUPABASE_KEY", "").strip()
SUPABASE_TABLE = os.environ.get("SUPABASE_TABLE", "registrations").strip()


def supabase_enabled():
    return bool(SUPABASE_URL and SUPABASE_KEY)


def supabase_available():
    if not supabase_enabled():
        return False
    try:
        resp = supabase_request("GET", f"{SUPABASE_TABLE}?select=no&limit=1")
    except Exception:
        return False
    return bool(resp is not None and getattr(resp, "status_code", 500) < 400)


def supabase_request(method, path, payload=None):
    if not supabase_enabled():
        return None
    url = f"{SUPABASE_URL.rstrip('/')}/rest/v1/{path}"
    headers = {
        "apikey": SUPABASE_KEY,
        "Authorization": f"Bearer {SUPABASE_KEY}",
        "Content-Type": "application/json",
        "Prefer": "return=minimal",
    }
    try:
        resp = requests.request(method, url, headers=headers, json=payload, timeout=15)
        if resp.status_code >= 400:
            return None
        return resp
    except requests.RequestException:
        return None


def sync_supabase_record(record):
    if not supabase_enabled() or not record:
        return None
    if "no" not in record:
        return None
    payload = {
        "no": str(record.get("no") or ""),
        "date": str(record.get("date") or ""),
        "name": str(record.get("name") or ""),
        "gender": str(record.get("gender") or ""),
        "dob": str(record.get("dob") or ""),
        "age": int(record.get("age") or 0),
        "mobile": str(record.get("mobile") or ""),
        "k_sar": str((record.get("k") or ["", "", "", ""])[0] or ""),
        "k_gam": str((record.get("k") or ["", "", "", ""])[1] or ""),
        "k_tal": str((record.get("k") or ["", "", "", ""])[2] or ""),
        "k_pin": str((record.get("k") or ["", "", "", ""])[3] or ""),
        "h_sar": str((record.get("h") or ["", "", "", ""])[0] or ""),
        "h_gam": str((record.get("h") or ["", "", "", ""])[1] or ""),
        "h_tal": str((record.get("h") or ["", "", "", ""])[2] or ""),
        "h_pin": str((record.get("h") or ["", "", "", ""])[3] or ""),
        "fee": int(record.get("fee") or 0),
        "ts": int(record.get("ts") or 0),
        "photo": str(record.get("photo") or ""),
        "aadhar": str(record.get("aadhar") or ""),
        "aadhar_no": str(record.get("aadhar_no") or ""),
        "role": str(record.get("role") or "operator"),
        "status": str(record.get("status") or "pending"),
        "approved_by": str(record.get("approved_by") or ""),
        "qr_code": str(record.get("qr_code") or ""),
    }
    try:
        supabase_request("POST", SUPABASE_TABLE, payload)
    except Exception:
        return None
    return payload


def supabase_row_to_record(row):
    if not isinstance(row, dict):
        return {}
    return {
        "no": row.get("no") or "",
        "date": row.get("date") or "",
        "name": row.get("name") or "",
        "gender": row.get("gender") or "bhai",
        "dob": row.get("dob") or "",
        "age": int(row.get("age") or 0),
        "mobile": row.get("mobile") or "",
        "k": [row.get("k_sar") or "", row.get("k_gam") or "", row.get("k_tal") or "", row.get("k_pin") or ""],
        "h": [row.get("h_sar") or "", row.get("h_gam") or "", row.get("h_tal") or "", row.get("h_pin") or ""],
        "fee": int(row.get("fee") or 0),
        "ts": int(row.get("ts") or 0),
        "photo": row.get("photo") or "",
        "aadhar": row.get("aadhar") or "",
        "aadhar_no": row.get("aadhar_no") or "",
        "role": row.get("role") or "operator",
        "status": row.get("status") or "pending",
        "approved_by": row.get("approved_by") or "",
        "qr_code": row.get("qr_code") or "",
    }


def list_supabase_records(q="", status=""):
    if not supabase_enabled() or not supabase_available():
        return []
    url = f"{SUPABASE_TABLE}?select=*"
    resp = supabase_request("GET", url)
    if resp is None:
        return []
    items = resp.json() if hasattr(resp, "json") else []
    qn = (q or "").lower().strip()
    statusn = (status or "").lower().strip()
    filtered = []
    for row in items:
        name = str((row.get("name") or "")).lower()
        no = str((row.get("no") or "")).lower()
        mobile = str((row.get("mobile") or "")).lower()
        item_status = str((row.get("status") or "pending")).lower()
        if qn and qn not in name and qn not in no and qn not in mobile:
            continue
        if statusn and item_status != statusn:
            continue
        filtered.append(supabase_row_to_record(row))
    return filtered


def get_supabase_record(no):
    if not supabase_enabled() or not supabase_available():
        return {}
    resp = supabase_request("GET", f"{SUPABASE_TABLE}?no=eq.{no}&select=*")
    if resp is None:
        return {}
    rows = resp.json() if hasattr(resp, "json") else []
    if not rows:
        return {}
    return supabase_row_to_record(rows[0])


# ------------------------- database -------------------------
def db():
    conn = sqlite3.connect(DB_PATH, timeout=15)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db() as c:
        c.execute(
            """CREATE TABLE IF NOT EXISTS registrations (
                no        TEXT PRIMARY KEY,
                date      TEXT NOT NULL,
                name      TEXT NOT NULL,
                gender    TEXT NOT NULL,
                dob       TEXT NOT NULL,
                age       INTEGER NOT NULL,
                mobile    TEXT NOT NULL,
                k_sar     TEXT, k_gam TEXT, k_tal TEXT, k_pin TEXT,
                h_sar     TEXT, h_gam TEXT, h_tal TEXT, h_pin TEXT,
                fee       INTEGER NOT NULL,
                ts        INTEGER NOT NULL,
                photo     TEXT, aadhar TEXT, aadhar_no TEXT,
                role      TEXT NOT NULL DEFAULT 'operator',
                status    TEXT NOT NULL DEFAULT 'pending',
                approved_by TEXT,
                qr_code   TEXT
            )"""
        )
        cols = {r["name"] for r in c.execute("PRAGMA table_info(registrations)").fetchall()}
        for col, ddl in (
            ("photo", "TEXT"),
            ("aadhar", "TEXT"),
            ("aadhar_no", "TEXT"),
            ("role", "TEXT NOT NULL DEFAULT 'operator'"),
            ("status", "TEXT NOT NULL DEFAULT 'pending'"),
            ("approved_by", "TEXT"),
            ("qr_code", "TEXT"),
        ):
            if col not in cols:
                c.execute(f"ALTER TABLE registrations ADD COLUMN {col} {ddl}")
        c.execute("CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
        c.execute("INSERT OR IGNORE INTO meta (key, value) VALUES ('counter', '0')")
        c.execute(
            """CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'operator',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )"""
        )


def hash_password(password: str) -> str:
    return hashlib.sha256(password.strip().encode("utf-8")).hexdigest()


def create_user_record(username: str, password: str, role: str = "operator"):
    username = (username or "").strip()
    password = (password or "").strip()
    role = (role or "operator").strip().lower()
    if not username or len(username) < 3:
        raise ValueError("યુઝરનેમ ઓછામાં ઓછી 3 અક્ષરનો હોવો જોઈએ.")
    if len(password) < 6:
        raise ValueError("પાસવર્ડ ઓછામાં ઓછી 6 અક્ષરનો હોવો જોઈએ.")
    if role not in VALID_ROLES:
        role = "operator"
    with db() as c:
        existing = c.execute("SELECT id, username, role FROM users WHERE username=?", (username,)).fetchone()
        if existing:
            return {"id": existing["id"], "username": existing["username"], "role": existing["role"]}
        c.execute(
            "INSERT INTO users (username, password, role) VALUES (?, ?, ?)",
            (username, hash_password(password), role),
        )
        row = c.execute("SELECT id, username, role FROM users WHERE username=?", (username,)).fetchone()
    return {"id": row["id"], "username": row["username"], "role": row["role"]}


def verify_user(username: str, password: str):
    username = (username or "").strip()
    password = (password or "").strip()
    if not username or not password:
        return None
    with db() as c:
        row = c.execute(
            "SELECT id, username, role FROM users WHERE username=? AND password=?",
            (username, hash_password(password)),
        ).fetchone()
    if not row:
        return None
    return {"id": row["id"], "username": row["username"], "role": row["role"]}


def get_counter():
    with db() as c:
        row = c.execute("SELECT value FROM meta WHERE key='counter'").fetchone()
    return int(row["value"])


def set_counter(n):
    with db() as c:
        c.execute("UPDATE meta SET value=? WHERE key='counter'", (str(n),))


def next_no():
    with _lock:
        n = get_counter() + 1
        set_counter(n)
        return f"NP-{date.today().year}-{n:04d}", n


def calc_age(dob_str: str):
    try:
        d = date.fromisoformat(dob_str)
    except (ValueError, TypeError):
        return None
    today = date.today()
    age = today.year - d.year - ((today.month, today.day) < (d.month, d.day))
    return age


def row_to_dict(r, with_blobs=False):
    d = {
        "no": r["no"],
        "date": r["date"],
        "name": r["name"],
        "gender": r["gender"],
        "dob": r["dob"],
        "age": r["age"],
        "mobile": r["mobile"],
        "k": [r["k_sar"], r["k_gam"], r["k_tal"], r["k_pin"]],
        "h": [r["h_sar"], r["h_gam"], r["h_tal"], r["h_pin"]],
        "fee": r["fee"],
        "ts": r["ts"],
        "aadhar_no": r["aadhar_no"] or "",
        "role": (r["role"] or "operator"),
        "status": (r["status"] or "pending"),
        "approved_by": r["approved_by"] or "",
        "qr_code": r["qr_code"] or "",
        "has_photo": bool(r["photo"]),
        "has_aadhar": bool(r["aadhar"]),
    }
    if with_blobs:
        d["photo"] = r["photo"] or ""
        d["aadhar"] = r["aadhar"] or ""
    return d


# ------------------------- settings -------------------------
def get_settings():
    with db() as c:
        row = c.execute("SELECT value FROM meta WHERE key='settings'").fetchone()
    if not row:
        return dict(DEFAULT_SETTINGS)
    try:
        s = json.loads(row["value"])
        return {**DEFAULT_SETTINGS, **s}
    except (ValueError, TypeError):
        return dict(DEFAULT_SETTINGS)


def set_settings(s):
    with db() as c:
        c.execute(
            "INSERT OR REPLACE INTO meta (key, value) VALUES ('settings', ?)",
            (json.dumps(s, ensure_ascii=False),),
        )


def logo_url():
    return "/logo.png" if (BASE_DIR / "navratri" / "logo.png").exists() else ""


def background_url():
    return "/background.png" if (BASE_DIR / "navratri" / "background.png").exists() else ""


def qr_url(value: str) -> str:
    payload = value.strip() or "navratri-pass"
    return f"https://api.qrserver.com/v1/create-qr-code/?size=180x180&data={__import__('urllib.parse').parse.quote(payload)}"


# ------------------------- admin pin -------------------------
def check_pin():
    if ADMIN_PIN is None or ADMIN_PIN == "":
        return True
    return request.headers.get("X-Admin-Pin", "").strip() == ADMIN_PIN


def pin_error():
    return jsonify({"ok": False, "error": "એડમિન PIN ખોટો છે."}), 401


# ------------------------- routes -------------------------
@app.route("/")
def index():
    return send_from_directory(STATIC_DIR, "index.html")


@app.route("/<path:filename>")
def static_files(filename):
    return send_from_directory(STATIC_DIR, filename)


@app.get("/api/health")
def health():
    return jsonify({"ok": True, "name": "Navratri Garba Pass 2026"})


@app.get("/api/next")
def api_next():
    return jsonify({"no": f"NP-{date.today().year}-{get_counter() + 1:04d}"})


@app.post("/api/users")
def api_create_user():
    data = request.get_json(silent=True) or {}
    username = str(data.get("username", "")).strip()
    password = str(data.get("password", "")).strip()
    role = str(data.get("role", "operator")).strip().lower()
    try:
        user = create_user_record(username, password, role)
        return jsonify({"ok": True, "user": user})
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400


@app.post("/api/login")
def api_login():
    data = request.get_json(silent=True) or {}
    username = str(data.get("username", "")).strip()
    password = str(data.get("password", "")).strip()
    user = verify_user(username, password)
    if not user:
        return jsonify({"ok": False, "error": "યુઝરનેમ અથવા પાસવર્ડ ખોટો છે."}), 401
    return jsonify({"ok": True, "user": user})


@app.post("/api/register")
def api_register():
    data = request.get_json(silent=True) or {}
    name = str(data.get("name", "")).strip()
    gender = str(data.get("gender", "")).strip()
    dob = str(data.get("dob", "")).strip()
    mobile = str(data.get("mobile", "")).strip()
    role = str(data.get("role", "operator")).strip().lower()
    if role not in VALID_ROLES:
        role = "operator"

    if not name:
        return jsonify({"ok": False, "error": "નામ ફરજિયાત છે."}), 400
    if gender not in GENDER_FEE:
        return jsonify({"ok": False, "error": "લિંગ પસંદ કરો (બહેન / ભાઈ)."}), 400
    if not MOBILE_RE.match(mobile):
        return jsonify({"ok": False, "error": "મોબાઈલ નંબર ૧૦ આંકડાનો ફરજિયાત છે."}), 400
    age = calc_age(dob)
    if age is None or age < 0:
        return jsonify({"ok": False, "error": "જન્મ તારીખ સાચી નથી."}), 400

    addr = data.get("k") or []
    haddr = data.get("h") or []
    k = [str(addr[i]).strip() if i < len(addr) else "" for i in range(4)]
    h = [str(haddr[i]).strip() if i < len(haddr) else "" for i in range(4)]

    aadhar_no = str(data.get("aadhar_no", "")).strip()
    if aadhar_no and not AADHAR_RE.match(aadhar_no):
        return jsonify({"ok": False, "error": "આધાર નંબર ૧૨ આંકડાનો ફરજિયાત છે."}), 400
    photo = str(data.get("photo") or "").strip()
    aadhar = str(data.get("aadhar") or "").strip()

    rec_no, _ = next_no()
    ts = int(__import__("time").time() * 1000)
    qr = qr_url(rec_no)
    with db() as c:
        c.execute(
            """INSERT INTO registrations
               (no, date, name, gender, dob, age, mobile,
                k_sar, k_gam, k_tal, k_pin, h_sar, h_gam, h_tal, h_pin, fee, ts,
                photo, aadhar, aadhar_no, role, status, approved_by, qr_code)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (rec_no, date.today().isoformat(), name, gender, dob, age, mobile,
             k[0], k[1], k[2], k[3], h[0], h[1], h[2], h[3], GENDER_FEE[gender], ts,
             photo, aadhar, aadhar_no, role, "pending", "", qr),
        )

    with db() as c:
        row = c.execute("SELECT * FROM registrations WHERE no=?", (rec_no,)).fetchone()
    record = row_to_dict(row, with_blobs=True)
    sync_supabase_record(record)
    return jsonify({"ok": True, "no": rec_no, "record": record})


@app.get("/api/registrations")
def api_list():
    q = request.args.get("q", "").strip().lower()
    status = request.args.get("status", "").strip().lower()
    with db() as c:
        if q:
            like = f"%{q}%"
            rows = c.execute(
                """SELECT * FROM registrations
                   WHERE (lower(name) LIKE ? OR mobile LIKE ? OR lower(no) LIKE ?)
                   AND (? = '' OR status = ?)
                   ORDER BY ts DESC""",
                (like, like, like, status, status),
            ).fetchall()
        else:
            if status:
                rows = c.execute(
                    "SELECT * FROM registrations WHERE status=? ORDER BY ts DESC",
                    (status,),
                ).fetchall()
            else:
                rows = c.execute("SELECT * FROM registrations ORDER BY ts DESC").fetchall()
    items = []
    for r in rows:
        d = row_to_dict(r)
        d["photo"] = r["photo"] or ""
        items.append(d)
    return jsonify(items)


@app.get("/api/registrations/<no>")
def api_get(no):

    with db() as c:
        row = c.execute("SELECT * FROM registrations WHERE no=?", (no,)).fetchone()
    if not row:
        return jsonify({"ok": False, "error": "પાસ મળ્યો નથી."}), 404
    return jsonify(row_to_dict(row, with_blobs=True))


@app.patch("/api/registrations/<no>")
def api_patch(no):
    if not check_pin():
        return pin_error()
    row = None
    with db() as c:
        row = c.execute("SELECT * FROM registrations WHERE no=?", (no,)).fetchone()
    if not row:
        return jsonify({"ok": False, "error": "પાસ મળ્યો નથી."}), 404

    data = request.get_json(silent=True) or {}

    status = str(data.get("status", row["status"] or "pending")).strip().lower()
    if status not in VALID_STATUSES:
        status = row["status"] or "pending"
    approved_by = str(data.get("approved_by", row["approved_by"] or "")).strip().lower()
    if status == "approved" and approved_by not in VALID_ROLES:
        approved_by = "authority"
    if status != "approved":
        approved_by = ""

    def field(name):
        return str(data.get(name, row[name] or "")).strip()

    name = field("name")
    mobile = field("mobile")
    dob = field("dob")
    rdate = field("date")
    if not name:
        return jsonify({"ok": False, "error": "નામ ફરજિયાત છે."}), 400
    if not mobile or not MOBILE_RE.match(mobile):
        return jsonify({"ok": False, "error": "મોબાઈલ નંબર ૧૦ આંકડાનો ફરજિયાત છે."}), 400
    age = calc_age(dob)
    if age is None or age < 0:
        return jsonify({"ok": False, "error": "જન્મ તારીખ સાચી નથી."}), 400

    aadhar_no = field("aadhar_no")
    if aadhar_no and not AADHAR_RE.match(aadhar_no):
        return jsonify({"ok": False, "error": "આધાર નંબર ૧૨ આંકડાનો ફરજિયાત છે."}), 400

    def addr(arr, cur):
        out = []
        for i in range(4):
            out.append(str(arr[i]).strip() if i < len(arr) else (cur[i] or ""))
        return out

    k = addr(data.get("k") or [], [row["k_sar"], row["k_gam"], row["k_tal"], row["k_pin"]])
    h = addr(data.get("h") or [], [row["h_sar"], row["h_gam"], row["h_tal"], row["h_pin"]])

    photo = data.get("photo")
    if photo is None:
        photo = row["photo"]
    aadhar = data.get("aadhar")
    if aadhar is None:
        aadhar = row["aadhar"]

    role = str(data.get("role", row["role"] or "operator")).strip().lower()
    if role not in VALID_ROLES:
        role = row["role"] or "operator"

    with db() as c:
        c.execute(
            """UPDATE registrations SET date=?, name=?, dob=?, age=?, mobile=?,
               k_sar=?, k_gam=?, k_tal=?, k_pin=?, h_sar=?, h_gam=?, h_tal=?, h_pin=?,
               photo=?, aadhar=?, aadhar_no=?, role=?, status=?, approved_by=?, qr_code=?
               WHERE no=?""",
            (rdate, name, dob, age, mobile,
             k[0], k[1], k[2], k[3], h[0], h[1], h[2], h[3],
             photo, aadhar, aadhar_no, role, status, approved_by, row["qr_code"] or qr_url(no), no),
        )
    with db() as c:
        row = c.execute("SELECT * FROM registrations WHERE no=?", (no,)).fetchone()
    record = row_to_dict(row, with_blobs=True)
    sync_supabase_record(record)
    return jsonify({"ok": True, "record": record})


@app.delete("/api/registrations/<no>")
def api_delete(no):
    if not check_pin():
        return pin_error()
    with db() as c:
        cur = c.execute("DELETE FROM registrations WHERE no=?", (no,))
    if cur.rowcount == 0:
        return jsonify({"ok": False, "error": "રજીસ્ટ્રેશન મળ્યું નથી."}), 404
    return jsonify({"ok": True})


@app.post("/api/clear")
def api_clear():
    if not check_pin():
        return pin_error()
    with _lock:
        with db() as c:
            c.execute("DELETE FROM registrations")
        set_counter(0)
    return jsonify({"ok": True})


@app.get("/api/settings")
def api_get_settings():
    s = get_settings()
    s["logo"] = s.get("logo") or logo_url()
    s["background"] = s.get("background") or background_url()
    return jsonify(s)


@app.post("/api/settings")
def api_save_settings():
    if not check_pin():
        return pin_error()
    data = request.get_json(silent=True) or {}
    s = get_settings()
    for key in ("title", "subtitle", "address", "date_from", "date_to", "background"):
        if key in data:
            s[key] = str(data[key]).strip()
    set_settings(s)
    s["logo"] = s.get("logo") or logo_url()
    s["background"] = s.get("background") or background_url()
    return jsonify({"ok": True, "settings": s})


@app.post("/api/logo")
def api_logo():
    if not check_pin():
        return pin_error()
    f = request.files.get("logo")
    if f is None:
        return jsonify({"ok": False, "error": "લોગો ફાઈલ પસંદ કરો."}), 400
    data = f.read()
    if not data:
        return jsonify({"ok": False, "error": "ખાલી ફાઈલ છે."}), 400
    if len(data) > 5 * 1024 * 1024:
        return jsonify({"ok": False, "error": "લોગો ફાઈલ ૫ MB થી નાની રાખો."}), 400
    (STATIC_DIR / "logo.png").write_bytes(data)
    return jsonify({"ok": True, "logo": "/logo.png"})


@app.post("/api/background")
def api_background():
    if not check_pin():
        return pin_error()
    f = request.files.get("background")
    if f is None:
        return jsonify({"ok": False, "error": "બેકગ્રાઉન્ડ ફાઈલ પસંદ કરો."}), 400
    data = f.read()
    if not data:
        return jsonify({"ok": False, "error": "ખાલી ફાઈલ છે."}), 400
    if len(data) > 5 * 1024 * 1024:
        return jsonify({"ok": False, "error": "બેકગ્રાઉન્ડ ફાઈલ ૫ MB થી નાની રાખો."}), 400
    (STATIC_DIR / "background.png").write_bytes(data)
    return jsonify({"ok": True, "background": "/background.png"})


@app.get("/api/stats")
def api_stats():
    with db() as c:
        total = c.execute("SELECT COUNT(*) AS n FROM registrations").fetchone()["n"]
        b = c.execute("SELECT COUNT(*) AS n, COALESCE(SUM(fee), 0) AS f FROM registrations WHERE gender='bahin'").fetchone()
        o = c.execute("SELECT COUNT(*) AS n, COALESCE(SUM(fee), 0) AS f FROM registrations WHERE gender='bhai'").fetchone()
    return jsonify({
        "total": total,
        "bahin": b["n"], "bahin_fee": b["f"],
        "bhai": o["n"], "bhai_fee": o["f"],
        "total_fee": b["f"] + o["f"],
    })


@app.get("/api/export.csv")
def api_export_csv():
    with db() as c:
        rows = c.execute("SELECT * FROM registrations ORDER BY ts DESC").fetchall()
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["નો", "તારીખ", "નામ", "લિંગ", "જન્મ તારીખ", "ઉમર", "મોબાઈલ",
                "કાયમી સરનામું", "ગામ", "તાલુકો", "પિન",
                "હાલ સરનામું", "ગામ", "તાલુકો", "પિન", "આધાર નંબર", "ફી"])
    for r in rows:
        w.writerow([r["no"], r["date"], r["name"], r["gender"], r["dob"], r["age"], r["mobile"],
                    r["k_sar"], r["k_gam"], r["k_tal"], r["k_pin"],
                    r["h_sar"], r["h_gam"], r["h_tal"], r["h_pin"],
                    r["aadhar_no"] or "", r["fee"]])
    data = "\ufeff" + buf.getvalue()
    return Response(
        data,
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=navratri_registrations.csv"},
    )


@app.get("/api/export.xlsx")
def api_export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    with db() as c:
        rows = c.execute("SELECT * FROM registrations ORDER BY ts DESC").fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Registrations"
    headers = ["નો.", "તારીખ", "નામ", "લિંગ", "જન્મ તારીખ", "ઉમર", "મોબાઈલ",
               "કાયમી સરનામું", "ગામ", "તાલુકો", "પિન",
               "હાલ સરનામું", "ગામ", "તાલુકો", "પિન", "આધાર નંબર", "ફી"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="8E1B2E")
        cell.alignment = Alignment(horizontal="center")

    gender_label = {"bahin": "બહેન", "bhai": "ભાઈ"}
    for r in rows:
        ws.append([r["no"], r["date"], r["name"], gender_label.get(r["gender"], r["gender"]),
                   r["dob"], r["age"], r["mobile"],
                   r["k_sar"], r["k_gam"], r["k_tal"], r["k_pin"],
                   r["h_sar"], r["h_gam"], r["h_tal"], r["h_pin"],
                   r["aadhar_no"] or "", r["fee"]])

    widths = [14, 12, 28, 8, 12, 6, 12, 24, 14, 14, 8, 24, 14, 14, 8, 16, 8]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=navratri_registrations.xlsx"},
    )


@app.get("/api/export.json")
def api_export_json():
    pin_ok = check_pin()
    if not pin_ok:
        return pin_error()
    with db() as c:
        rows = c.execute("SELECT * FROM registrations ORDER BY ts DESC").fetchall()
    payload = json.dumps({"counter": get_counter(), "ledger": [row_to_dict(r, with_blobs=True) for r in rows]},
                         ensure_ascii=False, indent=2)
    return Response(
        payload,
        mimetype="application/json; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=navratri_backup.json"},
    )


@app.post("/api/import")
def api_import():
    if not check_pin():
        return pin_error()
    data = request.get_json(silent=True) or {}
    ledger = data.get("ledger")
    if not isinstance(ledger, list):
        return jsonify({"ok": False, "error": "બેકઅપ ફાઈલ સાચી નથી."}), 400

    records = []
    for item in ledger:
        if not isinstance(item, dict) or not item.get("no") or not item.get("name"):
            continue
        age = item.get("age")
        if age is None:
            age = calc_age(str(item.get("dob", "")))
            if age is None:
                continue
        k = item.get("k") or [None] * 4
        h = item.get("h") or [None] * 4
        records.append((
            item.get("no"), item.get("date", date.today().isoformat()),
            item.get("name"), item.get("gender", "bhai"), item.get("dob", ""),
            age, item.get("mobile", ""),
            k[0] if k and len(k) > 0 else None,
            k[1] if k and len(k) > 1 else None,
            k[2] if k and len(k) > 2 else None,
            k[3] if k and len(k) > 3 else None,
            h[0] if h and len(h) > 0 else None,
            h[1] if h and len(h) > 1 else None,
            h[2] if h and len(h) > 2 else None,
            h[3] if h and len(h) > 3 else None,
            item.get("fee", 50), item.get("ts", 0) or 0,
            item.get("photo") or "", item.get("aadhar") or "",
            item.get("aadhar_no") or "",
        ))

    if not records:
        return jsonify({"ok": False, "error": "બેકઅપ ફાઈલમાં કોઈ લાયક રેકોર્ડ નથી."}), 400

    with _lock:
        with db() as c:
            c.execute("DELETE FROM registrations")
            c.executemany(
                """INSERT OR REPLACE INTO registrations
                   (no, date, name, gender, dob, age, mobile,
                    k_sar, k_gam, k_tal, k_pin, h_sar, h_gam, h_tal, h_pin, fee, ts,
                    photo, aadhar, aadhar_no)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                records,
            )
        nums = [int(r[0].rsplit("-", 1)[-1]) for r in records if r[0] and r[0].rsplit("-", 1)[-1].isdigit()]
        max_no = max(nums) if nums else 0
        given = int(str(data.get("counter", "0")))
        set_counter(max(max_no, given))

    return jsonify({"ok": True, "count": len(records)})


if __name__ == "__main__":
    init_db()
    host = os.environ.get("HOST", "0.0.0.0")
    port = int(os.environ.get("PORT", "5000"))
    print(f"Navratri Garba Pass server -> http://{host}:{port}")
    app.run(host=host, port=port, debug=False)