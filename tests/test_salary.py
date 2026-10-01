from io import BytesIO
import sqlite3

import pytest
from openpyxl import Workbook, load_workbook

import salary_server
from salary_service import calculate_salary


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(salary_server, "DB_PATH", tmp_path / "salary.db")
    salary_server.init_db()
    return salary_server.app.test_client()


def test_calculate_salary_matches_specification():
    result = calculate_salary(monthly_salary=20000, working_days=16, present_days=14)

    assert result["daily_salary"] == 1250
    assert result["calculated_salary"] == 17500
    assert result["monthly_salary"] == 20000


def test_salary_rejects_zero_working_days():
    with pytest.raises(ValueError, match="greater than zero"):
        calculate_salary(monthly_salary=20000, working_days=0, present_days=0)


def test_salary_api_generates_calendar_and_finalizes(client):
    department = client.post("/api/departments", json={"name": "Computer Science", "code": "CSE"})
    department_id = department.get_json()["id"]
    faculty = client.post(
        "/api/faculties",
        json={"employee_code": "EMP001", "name": "Test Faculty", "department_id": department_id, "monthly_salary": 20000, "weekdays": [0, 3]},
    )
    faculty_id = faculty.get_json()["id"]

    calendar_response = client.get(f"/api/attendance?faculty_id={faculty_id}&month=10&year=2026")
    calendar = calendar_response.get_json()
    assert calendar_response.status_code == 200
    assert len(calendar["attendance"]) == 31
    assert sum(row["is_working"] for row in calendar["attendance"]) == 9

    result = client.post("/api/salary/calculate", json={"faculty_id": faculty_id, "month": 10, "year": 2026})
    assert result.status_code == 200
    assert result.get_json()["salary"]["calculated_salary"] == "0.00"

    finalized = client.post(f"/api/salary/{faculty_id}/finalize", json={"month": 10, "year": 2026})
    assert finalized.status_code == 200
    assert finalized.get_json()["status"] == "finalized"


def test_departments_and_faculty_can_be_updated_and_soft_deleted(client):
    department = client.post("/api/departments", json={"name": "Physics", "code": "PHY"}).get_json()
    updated_department = client.patch(f"/api/departments/{department['id']}", json={"name": "Applied Physics", "code": "APP", "head_name": "Dr. Shah"})
    assert updated_department.status_code == 200

    faculty = client.post(
        "/api/faculties",
        json={"employee_code": "EMP002", "name": "Old Name", "department_id": department["id"], "monthly_salary": 10000, "joining_date": "2020-01-15", "relieving_date": "", "weekdays": [1, 4]},
    ).get_json()
    updated_faculty = client.patch(
        f"/api/faculties/{faculty['id']}",
        json={"employee_code": "EMP002", "name": "New Name", "department_id": department["id"], "monthly_salary": 12000, "joining_date": "2020-01-15", "relieving_date": "2026-09-30", "weekdays": [0, 2, 4]},
    )
    assert updated_faculty.status_code == 200
    updated_faculty = next(row for row in client.get("/api/faculties").get_json() if row["employee_code"] == "EMP002")
    assert updated_faculty["name"] == "New Name"
    assert updated_faculty["joining_date"] == "2020-01-15"
    assert updated_faculty["relieving_date"] == "2026-09-30"

    deleted = client.delete(f"/api/faculties/{faculty['id']}")
    assert deleted.status_code == 200
    assert not any(row["id"] == faculty["id"] for row in client.get("/api/faculties").get_json())
    assert client.delete(f"/api/departments/{department['id']}").status_code == 200


def test_faculty_edit_refreshes_existing_calendar_days(client):
    department_id = client.post("/api/departments", json={"name": "History", "code": "HIS"}).get_json()["id"]
    faculty = client.post(
        "/api/faculties",
        json={"employee_code": "EMP006", "name": "Calendar Faculty", "department_id": department_id, "monthly_salary": 18000, "weekdays": [0]},
    ).get_json()
    faculty_id = faculty["id"]
    client.get(f"/api/attendance?faculty_id={faculty_id}&month=10&year=2026")

    updated = client.patch(
        f"/api/faculties/{faculty_id}",
        json={"employee_code": "EMP006", "name": "Renamed Faculty", "department_id": department_id, "monthly_salary": 22000, "weekdays": [3], "email": "updated@example.com"},
    )
    assert updated.status_code == 200

    attendance = client.get(f"/api/attendance?faculty_id={faculty_id}&month=10&year=2026").get_json()
    assert attendance["faculty"]["name"] == "Renamed Faculty"
    assert attendance["faculty"]["monthly_salary"] == 22000
    assert attendance["faculty"]["email"] == "updated@example.com"
    assert sum(row["is_working"] for row in attendance["attendance"]) == 5


def test_special_permission_present_counts_for_attendance_and_pay_without_changing_scheduled_days(client):
    department_id = client.post("/api/departments", json={"name": "Design", "code": "DES"}).get_json()["id"]
    faculty_id = client.post(
        "/api/faculties",
        json={"employee_code": "EMP007", "name": "Special Day Faculty", "department_id": department_id, "monthly_salary": 20000, "weekdays": [0]},
    ).get_json()["id"]
    calendar = client.get(f"/api/attendance?faculty_id={faculty_id}&month=6&year=2026").get_json()
    sunday = next(row for row in calendar["attendance"] if row["attendance_date"] == "2026-06-07")
    assert sunday["is_working"] == 0

    updated = client.put(
        f"/api/attendance/{sunday['id']}",
        json={"is_working": False, "attendance_status": "special"},
    )

    assert updated.status_code == 200
    attendance = client.get(f"/api/attendance?faculty_id={faculty_id}&month=6&year=2026").get_json()
    special_day = next(row for row in attendance["attendance"] if row["attendance_date"] == "2026-06-07")
    assert special_day["is_working"] == 0
    assert special_day["attendance_status"] == "special"
    assert attendance["totals"]["working_days"] == 5
    assert attendance["totals"]["present_days"] == 1
    assert attendance["totals"]["regular_present_days"] == 0
    assert attendance["totals"]["special_present_days"] == 1

    payroll = client.post("/api/salary/calculate", json={"faculty_id": faculty_id, "month": 6, "year": 2026}).get_json()
    assert payroll["salary"]["present_days"] == "1"
    assert payroll["salary"]["special_present_days"] == "1"
    assert payroll["salary"]["calculated_salary"] == "4000.00"


def test_faculty_excel_export_contains_active_faculty(client):
    department_id = client.post("/api/departments", json={"name": "Mathematics", "code": "MAT"}).get_json()["id"]
    client.post(
        "/api/faculties",
        json={"employee_code": "EMP003", "name": "Export Faculty", "department_id": department_id, "monthly_salary": 25000, "joining_date": "2021-06-01", "relieving_date": "2026-09-30", "weekdays": [0, 3]},
    )

    response = client.get("/api/faculties/export.xlsx")
    workbook = load_workbook(BytesIO(response.data))
    sheet = workbook["Faculty Master"]

    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert sheet.max_row == 2
    assert sheet.cell(2, 1).value == "EMP003"
    assert sheet.cell(2, 2).value == "Export Faculty"
    assert sheet.cell(1, 7).value == "Joining Date"
    assert sheet.cell(1, 8).value == "Relieving Date"
    assert sheet.cell(2, 7).value == "2021-06-01"
    assert sheet.cell(2, 8).value == "2026-09-30"


def test_init_db_adds_relieving_date_to_existing_faculty_table(tmp_path, monkeypatch):
    database = tmp_path / "existing.db"
    with sqlite3.connect(database) as connection:
        connection.execute("CREATE TABLE faculties (id INTEGER PRIMARY KEY, joining_date TEXT NOT NULL DEFAULT '')")
    monkeypatch.setattr(salary_server, "DB_PATH", database)

    salary_server.init_db()

    with sqlite3.connect(database) as connection:
        columns = {row[1] for row in connection.execute("PRAGMA table_info(faculties)")}
    assert "relieving_date" in columns


def test_monthly_dashboard_summary_contains_all_faculty_and_total(client):
    department_id = client.post("/api/departments", json={"name": "English", "code": "ENG"}).get_json()["id"]
    for employee_code, name in [("EMP004", "First Faculty"), ("EMP005", "Second Faculty")]:
        client.post(
            "/api/faculties",
            json={"employee_code": employee_code, "name": name, "department_id": department_id, "monthly_salary": 20000, "weekdays": [0]},
        )

    response = client.get("/api/salary/monthly?month=10&year=2026")
    payload = response.get_json()

    assert response.status_code == 200
    assert [row["employee_code"] for row in payload["rows"]] == ["EMP004", "EMP005"]
    assert payload["rows"][0]["working_days"] == 4
    assert payload["rows"][0]["weekdays"] == [0]
    assert payload["total_payable"] == "0.00"

    report = client.get("/api/salary/monthly/export.xlsx?month=10&year=2026")
    workbook = load_workbook(BytesIO(report.data))
    sheet = workbook["Monthly Salary"]
    assert sheet.max_row == 4
    assert sheet.cell(4, 2).value == "TOTAL"


def test_faculty_views_sort_by_employee_code(client):
    department_id = client.post("/api/departments", json={"name": "Art", "code": "ART"}).get_json()["id"]
    for employee_code, name in [("Z9", "A Faculty"), ("A2", "Z Faculty")]:
        client.post(
            "/api/faculties",
            json={"employee_code": employee_code, "name": name, "department_id": department_id, "monthly_salary": 10000},
        )

    faculty_codes = [item["employee_code"] for item in client.get("/api/faculties").get_json()]
    summary_codes = [item["employee_code"] for item in client.get("/api/salary/monthly?month=6&year=2026").get_json()["rows"]]
    assert faculty_codes == ["A2", "Z9"]
    assert summary_codes == ["A2", "Z9"]


def visiting_salary_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "June -2026"
    sheet.cell(4, 1).value = "Month: - 01/06/2026 To 30/06/2026 Dept: B. Architecture"
    staff = [
        ("00086", "PROF. PREETY SHAH", [2, 3, 4, 5, 6, 16, 17, 18, 19, 22, 23, 24, 29, 30], 14, 14, 75000, 75000),
        ("00095", "PROF. NAMRATA VYAS", [], 0, 8, 26400, 0),
        ("01304", "PROF. HITEN J CHAVDA", [15, 16, 24, 29], 4, 14, 42000, 12000),
        ("01305", "PROF. HARDIK N TAMBOLI", [], 0, 8, 26400, 0),
    ]
    for row_number, (code, name, days, present, working, salary, payable) in enumerate(staff, start=8):
        sheet.cell(row_number, 1, code)
        sheet.cell(row_number, 2, name)
        for day in days:
            sheet.cell(row_number, day + 2, 1)
        sheet.cell(row_number, 33, present)
        sheet.cell(row_number, 34, working)
        sheet.cell(row_number, 35, salary)
        sheet.cell(row_number, 36, payable)
    sheet.cell(22, 2, "Preety Shah She is appointed with Fixed Salary (Monday, Tuedsday & Wednesday)")
    sheet.cell(23, 2, "Hiten Chavda Monday, Tuesday, Friday")
    sheet.cell(24, 2, "Hardik Tamboli Monday & Thursday")
    sheet.cell(26, 2, "2/3 Days In Week (Namrata Vyas)")
    sheet.cell(26, 3, "Tuesday & Thursday")
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def test_visiting_salary_sheet_import_preserves_june_totals_and_attendance(client):
    response = client.post(
        "/api/salary/visiting/import",
        data={"file": (visiting_salary_workbook(), "Visiting Salary Sheet.xlsx")},
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    assert response.get_json()["month"] == 6
    assert response.get_json()["year"] == 2026
    assert response.get_json()["total_payable"] == "87000.00"

    summary = client.get("/api/salary/monthly?month=6&year=2026").get_json()
    by_code = {row["employee_code"]: row for row in summary["rows"]}
    assert set(by_code) == {"00086", "00095", "01304", "01305"}
    assert (by_code["00086"]["working_days"], by_code["00086"]["present_days"], by_code["00086"]["final_salary"]) == (14, 14, "75000.00")
    assert (by_code["00095"]["working_days"], by_code["00095"]["present_days"], by_code["00095"]["final_salary"]) == (8, 0, "0.00")
    assert (by_code["01304"]["working_days"], by_code["01304"]["present_days"], by_code["01304"]["final_salary"]) == (14, 4, "12000.00")
    assert (by_code["01305"]["working_days"], by_code["01305"]["present_days"], by_code["01305"]["final_salary"]) == (8, 0, "0.00")

    report = client.get("/api/salary/monthly/export.xlsx?month=6&year=2026")
    workbook = load_workbook(BytesIO(report.data), data_only=True)
    sheet = workbook["Visiting Attendance"]
    assert sheet["A1"].value == "COLLEGE OF ARCHITECTURE"
    assert sheet["A3"].value == "Statement of Visiting Staff Attendance"
    assert sheet["C5"].value == "Date's of Attendance"
    assert sheet["AG5"].value == "Day's to\nattend"
    assert sheet["A8"].value == "00086"
    assert sheet["D8"].value == 1
    assert sheet["AG8"].value == 14
    assert sheet["AH8"].value == 14
    assert sheet["AJ12"].value == 87000
    assert sheet.page_setup.orientation == "landscape"

    repeated = client.post(
        "/api/salary/visiting/import",
        data={"file": (visiting_salary_workbook(), "Visiting Salary Sheet.xlsx")},
        content_type="multipart/form-data",
    )
    assert repeated.status_code == 200
    assert len(client.get("/api/faculties").get_json()) == 4
    assert client.get("/api/salary/visiting/latest-period").get_json() == {"month": 6, "year": 2026}
