import calendar
import json
import re
import sqlite3
from io import BytesIO
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from flask import Flask, jsonify, request, send_file, send_from_directory

from salary_service import calculate_salary


BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "salary"
DB_PATH = BASE_DIR / "salary_data.db"

app = Flask(__name__, static_folder=None)


def db():
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    return connection


def init_db():
    with db() as connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS departments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE,
                code TEXT NOT NULL UNIQUE,
                head_name TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT 'active'
            );
            CREATE TABLE IF NOT EXISTS faculties (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                employee_code TEXT NOT NULL UNIQUE,
                name TEXT NOT NULL,
                department_id INTEGER NOT NULL,
                designation TEXT NOT NULL DEFAULT '',
                faculty_type TEXT NOT NULL DEFAULT 'custom',
                monthly_salary NUMERIC NOT NULL CHECK(monthly_salary >= 0),
                joining_date TEXT NOT NULL DEFAULT '',
                relieving_date TEXT NOT NULL DEFAULT '',
                email TEXT NOT NULL DEFAULT '',
                mobile TEXT NOT NULL DEFAULT '',
                weekdays TEXT NOT NULL DEFAULT '[0, 1, 2, 3, 4]',
                status TEXT NOT NULL DEFAULT 'active',
                FOREIGN KEY (department_id) REFERENCES departments(id)
            );
            CREATE TABLE IF NOT EXISTS monthly_attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                faculty_id INTEGER NOT NULL,
                attendance_date TEXT NOT NULL,
                is_working INTEGER NOT NULL DEFAULT 0,
                attendance_status TEXT NOT NULL DEFAULT '',
                leave_type TEXT NOT NULL DEFAULT '',
                remarks TEXT NOT NULL DEFAULT '',
                UNIQUE(faculty_id, attendance_date),
                FOREIGN KEY (faculty_id) REFERENCES faculties(id)
            );
            CREATE TABLE IF NOT EXISTS salary_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                faculty_id INTEGER NOT NULL,
                month INTEGER NOT NULL,
                year INTEGER NOT NULL,
                snapshot TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'draft',
                UNIQUE(faculty_id, month, year),
                FOREIGN KEY (faculty_id) REFERENCES faculties(id)
            );
            CREATE TABLE IF NOT EXISTS audit_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                action TEXT NOT NULL,
                module TEXT NOT NULL,
                record_id TEXT NOT NULL,
                old_data TEXT NOT NULL DEFAULT '',
                new_data TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            """
        )
        faculty_columns = {row["name"] for row in connection.execute("PRAGMA table_info(faculties)")}
        if "relieving_date" not in faculty_columns:
            connection.execute("ALTER TABLE faculties ADD COLUMN relieving_date TEXT NOT NULL DEFAULT ''")


def row_dict(row):
    return dict(row) if row else None


def faculty_dict(row):
    item = row_dict(row)
    item["weekdays"] = json.loads(item["weekdays"])
    item["monthly_salary"] = float(item["monthly_salary"])
    return item


def error(message, status=400):
    return jsonify({"ok": False, "error": message}), status


def request_int(value, default):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def ensure_calendar(faculty_id, month, year):
    with db() as connection:
        faculty = connection.execute("SELECT * FROM faculties WHERE id=?", (faculty_id,)).fetchone()
        if not faculty:
            return None, []
        weekdays = set(json.loads(faculty["weekdays"]))
        last_day = calendar.monthrange(year, month)[1]
        for day in range(1, last_day + 1):
            current = date(year, month, day)
            connection.execute(
                "INSERT OR IGNORE INTO monthly_attendance (faculty_id, attendance_date, is_working) VALUES (?, ?, ?)",
                (faculty_id, current.isoformat(), int(current.weekday() in weekdays)),
            )
        rows = connection.execute(
            "SELECT * FROM monthly_attendance WHERE faculty_id=? AND attendance_date LIKE ? ORDER BY attendance_date",
            (faculty_id, f"{year:04d}-{month:02d}-%"),
        ).fetchall()
    return faculty, [row_dict(row) for row in rows]


def sync_existing_calendar_weekdays(faculty_id, weekdays):
    weekdays = set(weekdays)
    with db() as connection:
        rows = connection.execute("SELECT id, attendance_date FROM monthly_attendance WHERE faculty_id=?", (faculty_id,)).fetchall()
        for row in rows:
            is_working = int(date.fromisoformat(row["attendance_date"]).weekday() in weekdays)
            connection.execute(
                "UPDATE monthly_attendance SET is_working=?, attendance_status=CASE WHEN ?=0 THEN '' ELSE attendance_status END, leave_type=CASE WHEN ?=0 THEN '' ELSE leave_type END WHERE id=?",
                (is_working, is_working, is_working, row["id"]),
            )


def attendance_totals(rows):
    working = [row for row in rows if row["is_working"]]
    regular_present = sum(row["attendance_status"] == "present" for row in working)
    special_present = sum(not row["is_working"] and row["attendance_status"] == "special" for row in rows)
    paid_leave = sum(row["attendance_status"] == "leave" and row["leave_type"] == "paid" for row in working)
    unpaid_leave = sum(row["attendance_status"] == "leave" and row["leave_type"] == "unpaid" for row in working)
    absent = sum(row["attendance_status"] == "absent" for row in working)
    return {"working_days": len(working), "present_days": regular_present + special_present, "regular_present_days": regular_present, "special_present_days": special_present, "absent_days": absent, "paid_leave": paid_leave, "unpaid_leave": unpaid_leave}


def visiting_sheet_weekdays(sheet, name, attendance_dates):
    normalize = lambda value: re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()
    name_parts = [part for part in normalize(name).split() if part not in {"prof", "professor"}]
    identity = {name_parts[0], name_parts[-1]} if name_parts else set()
    weekday_patterns = (
        (0, r"\bmon(?:day)?\b"),
        (1, r"\b(?:tue(?:sday)?|tuedsday)\b"),
        (2, r"\bwed(?:nesday)?\b"),
        (3, r"\bthu(?:rsday)?\b"),
        (4, r"\bfri(?:day)?\b"),
        (5, r"\bsat(?:urday)?\b"),
        (6, r"\bsun(?:day)?\b"),
    )
    for row in sheet.iter_rows(min_row=18, values_only=True):
        note = " ".join(str(value) for value in row if value is not None)
        normalized_note = normalize(note)
        if identity and all(part in normalized_note.split() for part in identity):
            weekdays = [day for day, pattern in weekday_patterns if re.search(pattern, normalized_note)]
            if weekdays:
                return weekdays
    if attendance_dates:
        return sorted({date.fromisoformat(item).weekday() for item in attendance_dates})
    return [0, 1, 2, 3, 4]


def parse_visiting_salary_sheet(workbook):
    from openpyxl.utils.exceptions import InvalidFileException

    sheet = next((item for item in workbook.worksheets if re.search(r"june\s*-?\s*2026", item.title, re.I)), None)
    if sheet is None:
        raise ValueError("Workbook must contain a June 2026 sheet")
    heading = str(sheet.cell(4, 1).value or "")
    period_match = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", heading)
    if not period_match:
        raise ValueError("Could not read the month and year from the June sheet")
    period_start = datetime.strptime("/".join(period_match.groups()), "%d/%m/%Y").date()
    month, year = period_start.month, period_start.year
    department_match = re.search(r"Dept:\s*([^\n]+)", heading, re.I)
    department_name = department_match.group(1).strip() if department_match else "Visiting Staff"
    department_name = re.split(r"\bDate\s*:", department_name, maxsplit=1, flags=re.I)[0].strip()
    department_name = re.sub(r"\s+", " ", department_name)
    department_name = re.sub(r"^[A-Z]\.\s*", "", department_name).strip()

    staff = []
    for row_number in range(8, sheet.max_row + 1):
        row = [sheet.cell(row_number, column).value for column in range(1, 37)]
        employee_code, name = str(row[0] or "").strip(), str(row[1] or "").strip()
        if not employee_code or not name or not re.fullmatch(r"[A-Za-z0-9-]+", employee_code):
            continue
        try:
            present_days, working_days = int(row[32]), int(row[33])
            monthly_salary, sheet_payable = Decimal(str(row[34])), Decimal(str(row[35]))
        except (TypeError, ValueError, ArithmeticError) as exc:
            raise ValueError(f"Invalid salary or attendance values for {name}") from exc
        if present_days < 0 or working_days <= 0 or present_days > working_days or monthly_salary < 0:
            raise ValueError(f"Invalid working or attendance day counts for {name}")

        attendance_dates = []
        for day, marker in enumerate(row[2:32], start=1):
            if str(marker or "").strip() in {"1", "1.0"}:
                try:
                    attendance_dates.append(date(year, month, day).isoformat())
                except ValueError as exc:
                    raise ValueError(f"Attendance date {day} is outside {month}/{year}") from exc
        if len(attendance_dates) != present_days:
            raise ValueError(f"Attendance marks do not match the sheet total for {name}")
        weekdays = visiting_sheet_weekdays(sheet, name, attendance_dates)
        result = calculate_salary(
            monthly_salary=monthly_salary,
            working_days=working_days,
            present_days=present_days,
        )
        if Decimal(str(result["final_salary"])) != sheet_payable:
            raise ValueError(f"Calculated pay does not match the sheet for {name}")
        staff.append({
            "employee_code": employee_code,
            "name": name,
            "monthly_salary": monthly_salary,
            "sheet_payable": sheet_payable,
            "present_days": present_days,
            "working_days": working_days,
            "attendance_dates": attendance_dates,
            "weekdays": weekdays,
            "calculation": result,
        })
    if not staff:
        raise ValueError("No visiting staff rows were found in the June sheet")
    if len({item["employee_code"] for item in staff}) != len(staff):
        raise ValueError("The June sheet contains duplicate employee codes")
    return month, year, department_name, staff


def import_visiting_salary_data(month, year, department_name, staff):
    last_day = calendar.monthrange(year, month)[1]
    with db() as connection:
        department = connection.execute(
            "SELECT id FROM departments WHERE lower(name)=lower(?) AND status='active'",
            (department_name,),
        ).fetchone()
        if department:
            department_id = department["id"]
        else:
            code_base = re.sub(r"[^A-Z0-9]", "", department_name.upper())[:8] or "VISITING"
            department_code, suffix = code_base, 1
            while connection.execute("SELECT 1 FROM departments WHERE code=?", (department_code,)).fetchone():
                suffix += 1
                department_code = f"{code_base[:6]}{suffix}"
            cursor = connection.execute(
                "INSERT INTO departments (name, code) VALUES (?, ?)",
                (department_name, department_code),
            )
            department_id = cursor.lastrowid

        faculty_ids = {}
        for item in staff:
            existing = connection.execute(
                "SELECT id, name FROM faculties WHERE employee_code=?",
                (item["employee_code"],),
            ).fetchone()
            if existing and re.sub(r"[^a-z0-9]", "", existing["name"].lower()) != re.sub(r"[^a-z0-9]", "", item["name"].lower()):
                raise ValueError(f"Employee code {item['employee_code']} belongs to another faculty member")
            faculty_id = existing["id"] if existing else None
            if faculty_id and connection.execute(
                "SELECT 1 FROM salary_records WHERE faculty_id=? AND month=? AND year=? AND status='finalized'",
                (faculty_id, month, year),
            ).fetchone():
                raise ValueError(f"{item['name']} has finalized salary for this period")
            faculty_ids[item["employee_code"]] = faculty_id

        for item in staff:
            weekdays = json.dumps(item["weekdays"])
            faculty_id = faculty_ids[item["employee_code"]]
            if faculty_id:
                connection.execute(
                    "UPDATE faculties SET name=?, department_id=?, designation='Visiting Faculty', faculty_type='visiting', monthly_salary=?, weekdays=? WHERE id=?",
                    (item["name"], department_id, str(item["monthly_salary"]), weekdays, faculty_id),
                )
            else:
                cursor = connection.execute(
                    "INSERT INTO faculties (employee_code, name, department_id, designation, faculty_type, monthly_salary, weekdays) VALUES (?, ?, ?, 'Visiting Faculty', 'visiting', ?, ?)",
                    (item["employee_code"], item["name"], department_id, str(item["monthly_salary"]), weekdays),
                )
                faculty_id = cursor.lastrowid
                faculty_ids[item["employee_code"]] = faculty_id

            present_dates = set(item["attendance_dates"])
            working_dates = {
                date(year, month, day).isoformat()
                for day in range(1, last_day + 1)
                if date(year, month, day).weekday() in item["weekdays"]
            }
            working_dates.update(present_dates)
            if len(working_dates) > item["working_days"]:
                removable = sorted(working_dates - present_dates, reverse=True)
                while len(working_dates) > item["working_days"] and removable:
                    working_dates.remove(removable.pop(0))
            if len(working_dates) < item["working_days"]:
                for day in range(1, last_day + 1):
                    attendance_date = date(year, month, day).isoformat()
                    if attendance_date not in working_dates:
                        working_dates.add(attendance_date)
                        if len(working_dates) == item["working_days"]:
                            break
            if len(working_dates) != item["working_days"]:
                raise ValueError(f"Could not build the attendance calendar for {item['name']}")

            for day in range(1, last_day + 1):
                attendance_date = date(year, month, day).isoformat()
                is_working = int(attendance_date in working_dates)
                status = "present" if attendance_date in present_dates else "absent" if is_working else ""
                connection.execute(
                    "INSERT INTO monthly_attendance (faculty_id, attendance_date, is_working, attendance_status) VALUES (?, ?, ?, ?) ON CONFLICT(faculty_id, attendance_date) DO UPDATE SET is_working=excluded.is_working, attendance_status=excluded.attendance_status, leave_type='', remarks=''",
                    (faculty_id, attendance_date, is_working, status),
                )
            snapshot = {key: str(value) for key, value in item["calculation"].items()}
            connection.execute(
                "INSERT INTO salary_records (faculty_id, month, year, snapshot, status) VALUES (?, ?, ?, ?, 'draft') ON CONFLICT(faculty_id, month, year) DO UPDATE SET snapshot=excluded.snapshot, status='draft'",
                (faculty_id, month, year, json.dumps(snapshot)),
            )
            connection.execute(
                "INSERT INTO audit_logs (action, module, record_id, new_data) VALUES (?, 'salary', ?, ?)",
                ("Visiting salary sheet imported", str(faculty_id), json.dumps({"month": month, "year": year, "payable": str(item["sheet_payable"])})),
            )
    return {
        "month": month,
        "year": year,
        "staff": [{key: item[key] for key in ("employee_code", "name", "monthly_salary", "working_days", "present_days", "sheet_payable")} for item in staff],
        "total_payable": f"{sum(item['sheet_payable'] for item in staff):.2f}",
    }


def monthly_salary_rows(month, year):
    with db() as connection:
        faculty_rows = connection.execute(
            "SELECT f.*, d.name AS department_name FROM faculties f JOIN departments d ON d.id=f.department_id WHERE f.status='active' ORDER BY f.employee_code COLLATE NOCASE"
        ).fetchall()
    rows = []
    for faculty in faculty_rows:
        _, attendance = ensure_calendar(faculty["id"], month, year)
        totals = attendance_totals(attendance)
        try:
            salary = calculate_salary(
                monthly_salary=faculty["monthly_salary"],
                working_days=totals["working_days"],
                present_days=totals["regular_present_days"],
                paid_leave=totals["paid_leave"],
                unpaid_leave=totals["unpaid_leave"],
                special_present_days=totals["special_present_days"],
            )
            salary_error = ""
        except ValueError as exc:
            salary = {"monthly_salary": faculty["monthly_salary"], "calculated_salary": 0, "final_salary": 0}
            salary_error = str(exc)
        rows.append({
            "faculty_id": faculty["id"],
            "employee_code": faculty["employee_code"],
            "name": faculty["name"],
            "department": faculty["department_name"],
            "faculty_type": faculty["faculty_type"],
            "weekdays": json.loads(faculty["weekdays"]),
            "monthly_salary": str(salary["monthly_salary"]),
            "working_days": totals["working_days"],
            "present_days": totals["present_days"],
            "absent_days": totals["absent_days"],
            "paid_leave": totals["paid_leave"],
            "unpaid_leave": totals["unpaid_leave"],
            "calculated_salary": str(salary["calculated_salary"]),
            "final_salary": str(salary["final_salary"]),
            "error": salary_error,
        })
    return rows


def visiting_attendance_workbook(month, year, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    days_in_month = calendar.monthrange(year, month)[1]
    first_date, last_date = date(year, month, 1), date(year, month, days_in_month)
    day_start, day_end = 3, days_in_month + 2
    days_to_attend_col, working_days_col = days_in_month + 3, days_in_month + 4
    salary_col, payable_col = days_in_month + 5, days_in_month + 6
    final_col = payable_col
    departments = sorted({row["department"] for row in rows})
    department_label = departments[0] if len(departments) == 1 else "All Departments"
    report_title = "Statement of Visiting Staff Attendance" if all(row["faculty_type"] == "visiting" for row in rows) else "Statement of Faculty Attendance"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Visiting Attendance"
    sheet.sheet_view.showGridLines = False

    def merged_cell(row, start_col, end_col, value, *, font=None, alignment=None):
        sheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
        cell = sheet.cell(row, start_col, value)
        if font:
            cell.font = font
        if alignment:
            cell.alignment = alignment
        return cell

    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    title_font = Font(name="Times New Roman", size=14, bold=True, underline="single")
    subtitle_font = Font(name="Times New Roman", size=12, underline="single")
    report_font = Font(name="Times New Roman", size=12, bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBE7")
    header_font = Font(name="Arial", size=9, bold=True)
    thin = Side(style="thin", color="404B4A")
    grid_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    merged_cell(1, 1, final_col, "COLLEGE OF ARCHITECTURE", font=title_font, alignment=centered)
    merged_cell(2, 1, final_col, "Sardar Vallabhbhai Patel Institute of Technology (S.V.I.T.) Vasad.", font=subtitle_font, alignment=centered)
    merged_cell(3, 1, final_col, report_title, font=report_font, alignment=centered)
    merged_cell(4, 1, 2, f"Month: - {first_date:%d/%m/%Y} To {last_date:%d/%m/%Y}", alignment=Alignment(vertical="center"))
    merged_cell(4, day_start, day_end, f"Dept: {department_label}", font=Font(bold=True), alignment=centered)
    merged_cell(4, days_to_attend_col, payable_col, f"Date: {last_date:%d/%m/%Y}", alignment=Alignment(horizontal="right", vertical="center"))

    sheet.merge_cells(start_row=5, start_column=1, end_row=7, end_column=1)
    sheet.merge_cells(start_row=5, start_column=2, end_row=7, end_column=2)
    sheet.cell(5, 1, "Emp.\nCode")
    sheet.cell(5, 2, "Name of\nVisiting Staff")
    sheet.merge_cells(start_row=5, start_column=day_start, end_row=5, end_column=day_end)
    sheet.cell(5, day_start, "Date's of Attendance")
    headers = (
        (days_to_attend_col, "Day's to\nattend"),
        (working_days_col, "Day's attended\nduring the month"),
        (salary_col, "Salary"),
        (payable_col, "Total salary\nfor pay"),
    )
    for column, label in headers:
        sheet.merge_cells(start_row=5, start_column=column, end_row=7, end_column=column)
        sheet.cell(5, column, label)
    for day in range(1, days_in_month + 1):
        column = day_start + day - 1
        current = date(year, month, day)
        sheet.cell(6, column, day)
        sheet.cell(7, column, current.strftime("%a"))

    for row_number in range(5, 8):
        for column in range(1, final_col + 1):
            cell = sheet.cell(row_number, column)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = centered
            cell.border = grid_border

    with db() as connection:
        for row_number, item in enumerate(rows, start=8):
            attendance_rows = connection.execute(
                "SELECT attendance_date, is_working, attendance_status FROM monthly_attendance WHERE faculty_id=? AND attendance_date LIKE ?",
                (item["faculty_id"], f"{year:04d}-{month:02d}-%"),
            ).fetchall()
            attendance_by_day = {int(entry["attendance_date"][-2:]): entry for entry in attendance_rows}
            values = {
                1: item["employee_code"],
                2: item["name"],
                days_to_attend_col: item["present_days"],
                working_days_col: item["working_days"],
                salary_col: float(item["monthly_salary"]),
                payable_col: float(item["final_salary"]),
            }
            for column, value in values.items():
                cell = sheet.cell(row_number, column, value)
                cell.border = grid_border
                cell.alignment = Alignment(horizontal="left" if column <= 2 else "center", vertical="center", wrap_text=True)
            for day in range(1, days_in_month + 1):
                entry = attendance_by_day.get(day)
                status = entry["attendance_status"] if entry else ""
                is_working = bool(entry["is_working"]) if entry else False
                cell = sheet.cell(row_number, day_start + day - 1, 1 if status in {"present", "special"} else None)
                cell.alignment = centered
                cell.border = grid_border
                if status == "special":
                    cell.fill = PatternFill("solid", fgColor="FFF0BF")
                elif not is_working:
                    cell.fill = PatternFill("solid", fgColor="D9D9D9")

    total_row = len(rows) + 8
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=payable_col - 1)
    total_label = f"TOTAL {department_label.upper()} DEPT. VISITING STAFF SALARY RS."
    sheet.cell(total_row, 1, total_label)
    sheet.cell(total_row, payable_col, sum(float(row["final_salary"]) for row in rows))
    for column in range(1, final_col + 1):
        cell = sheet.cell(total_row, column)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = grid_border
    sheet.cell(total_row, 1).alignment = Alignment(vertical="center")
    sheet.cell(total_row, payable_col).number_format = '#,##0.00'

    signature_row = total_row + 4
    merged_cell(signature_row, 2, min(5, final_col), "Authorized Signatory\nCollege of Architecture\nSVIT - Vasad")
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 30
    for column in range(day_start, day_end + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 4.5
    for column, width in ((days_to_attend_col, 12), (working_days_col, 15), (salary_col, 13), (payable_col, 15)):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row_number, height in ((1, 24), (2, 22), (3, 22), (4, 24), (5, 36), (6, 22), (7, 26)):
        sheet.row_dimensions[row_number].height = height
    sheet.freeze_panes = "C8"
    sheet.auto_filter.ref = f"A5:{get_column_letter(final_col)}{total_row - 1}"
    sheet.print_area = f"A1:{get_column_letter(final_col)}{signature_row + 2}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.2, right=0.2, top=0.35, bottom=0.35, header=0.15, footer=0.15)
    return workbook


@app.post("/api/salary/visiting/import")
def import_visiting_salary_sheet():
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    from zipfile import BadZipFile

    upload = request.files.get("file")
    if not upload or not upload.filename:
        return error("Choose a visiting salary workbook")
    try:
        workbook = load_workbook(BytesIO(upload.read()), data_only=True, read_only=True)
        month, year, department_name, staff = parse_visiting_salary_sheet(workbook)
        result = import_visiting_salary_data(month, year, department_name, staff)
    except ValueError as exc:
        return error(str(exc), 409 if "finalized" in str(exc) or "belongs to another" in str(exc) else 400)
    except (InvalidFileException, BadZipFile, OSError, KeyError):
        return error("Could not read the uploaded Excel workbook")
    return jsonify(result)


@app.get("/api/salary/visiting/latest-period")
def latest_visiting_salary_period():
    with db() as connection:
        row = connection.execute(
            "SELECT r.month, r.year FROM salary_records r JOIN faculties f ON f.id=r.faculty_id WHERE f.faculty_type='visiting' ORDER BY r.year DESC, r.month DESC LIMIT 1"
        ).fetchone()
    return jsonify({"month": row["month"], "year": row["year"]} if row else {"month": None, "year": None})


@app.route("/")
def index():
    return send_from_directory(STATIC_DIR, "index.html")


@app.get("/api/health")
def health():
    return jsonify({"ok": True, "name": "Faculty Salary Management"})


@app.get("/api/departments")
def departments():
    with db() as connection:
        rows = connection.execute("SELECT * FROM departments WHERE status='active' ORDER BY name").fetchall()
    return jsonify([row_dict(row) for row in rows])


@app.post("/api/departments")
def create_department():
    data = request.get_json(silent=True) or {}
    name, code = str(data.get("name", "")).strip(), str(data.get("code", "")).strip().upper()
    if not name or not code:
        return error("Department name and code are required")
    try:
        with db() as connection:
            cursor = connection.execute("INSERT INTO departments (name, code, head_name) VALUES (?, ?, ?)", (name, code, str(data.get("head_name", "")).strip()))
        return jsonify({"ok": True, "id": cursor.lastrowid}), 201
    except sqlite3.IntegrityError:
        return error("Department name or code already exists", 409)


@app.patch("/api/departments/<int:department_id>")
def update_department(department_id):
    data = request.get_json(silent=True) or {}
    name, code = str(data.get("name", "")).strip(), str(data.get("code", "")).strip().upper()
    if not name or not code:
        return error("Department name and code are required")
    try:
        with db() as connection:
            cursor = connection.execute("UPDATE departments SET name=?, code=?, head_name=? WHERE id=? AND status='active'", (name, code, str(data.get("head_name", "")).strip(), department_id))
        if cursor.rowcount == 0:
            return error("Department not found", 404)
        return jsonify({"ok": True})
    except sqlite3.IntegrityError:
        return error("Department name or code already exists", 409)


@app.delete("/api/departments/<int:department_id>")
def delete_department(department_id):
    with db() as connection:
        faculty_count = connection.execute("SELECT COUNT(*) FROM faculties WHERE department_id=? AND status='active'", (department_id,)).fetchone()[0]
        if faculty_count:
            return error("Move or delete the department faculty first", 409)
        cursor = connection.execute("UPDATE departments SET status='inactive' WHERE id=? AND status='active'", (department_id,))
    if cursor.rowcount == 0:
        return error("Department not found", 404)
    return jsonify({"ok": True})


@app.get("/api/faculties")
def faculties():
    with db() as connection:
        rows = connection.execute("SELECT f.*, d.name AS department_name FROM faculties f JOIN departments d ON d.id=f.department_id WHERE f.status='active' ORDER BY f.employee_code COLLATE NOCASE").fetchall()
    return jsonify([faculty_dict(row) for row in rows])


@app.get("/api/faculties/export.xlsx")
def export_faculties():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    with db() as connection:
        rows = connection.execute(
            "SELECT f.*, d.name AS department_name FROM faculties f JOIN departments d ON d.id=f.department_id WHERE f.status='active' ORDER BY f.employee_code COLLATE NOCASE"
        ).fetchall()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Faculty Master"
    headers = ["Employee Code", "Faculty Name", "Department", "Designation", "Faculty Type", "Monthly Salary", "Joining Date", "Relieving Date", "Email", "Mobile", "Working Weekdays", "Status"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="116B67")
    day_names = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    for row in rows:
        weekdays = ", ".join(day_names[index] for index in json.loads(row["weekdays"]))
        sheet.append([
            row["employee_code"], row["name"], row["department_name"], row["designation"],
            row["faculty_type"], float(row["monthly_salary"]), row["joining_date"], row["relieving_date"], row["email"],
            row["mobile"], weekdays, row["status"],
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 30)
        sheet.column_dimensions[column[0].column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="Faculty_Master.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/salary/monthly")
def monthly_salary():
    month = request_int(request.args.get("month"), date.today().month)
    year = request_int(request.args.get("year"), date.today().year)
    if month not in range(1, 13) or year < 2000:
        return error("Month and year are required")
    rows = monthly_salary_rows(month, year)
    return jsonify({"month": month, "year": year, "rows": rows, "total_payable": f"{sum(float(row['final_salary']) for row in rows):.2f}"})


@app.get("/api/salary/monthly/export.xlsx")
def export_monthly_salary():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    month = request_int(request.args.get("month"), date.today().month)
    year = request_int(request.args.get("year"), date.today().year)
    if month not in range(1, 13) or year < 2000:
        return error("Month and year are required")
    rows = monthly_salary_rows(month, year)
    if any(row["faculty_type"] == "visiting" for row in rows):
        workbook = visiting_attendance_workbook(month, year, rows)
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=f"Visiting_Salary_{year}_{month:02d}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monthly Salary"
    headers = ["Employee Code", "Faculty Name", "Department", "Monthly Salary", "Working Days", "Present Days", "Absent Days", "Paid Leave", "Unpaid Leave", "Calculated Salary", "Total Payable"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="116B67")
    for row in rows:
        sheet.append([row["employee_code"], row["name"], row["department"], float(row["monthly_salary"]), row["working_days"], row["present_days"], row["absent_days"], row["paid_leave"], row["unpaid_leave"], float(row["calculated_salary"]), float(row["final_salary"])])
    sheet.append(["", "TOTAL", "", "", "", "", "", "", "", "", sum(float(row["final_salary"]) for row in rows)])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:K{len(rows) + 1}"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 28)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"Salary_Report_{year}_{month:02d}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.post("/api/faculties")
def create_faculty():
    data = request.get_json(silent=True) or {}
    required = ["employee_code", "name", "department_id", "monthly_salary"]
    if any(str(data.get(field, "")).strip() == "" for field in required):
        return error("Employee code, name, department, and monthly salary are required")
    weekdays = data.get("weekdays", [0, 1, 2, 3, 4])
    try:
        weekdays = sorted({int(day) for day in weekdays if 0 <= int(day) <= 6})
        salary = float(data["monthly_salary"])
        if salary < 0:
            raise ValueError
        with db() as connection:
            cursor = connection.execute(
                "INSERT INTO faculties (employee_code, name, department_id, designation, faculty_type, monthly_salary, joining_date, relieving_date, email, mobile, weekdays) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (str(data["employee_code"]).strip().upper(), str(data["name"]).strip(), int(data["department_id"]), str(data.get("designation", "")).strip(), str(data.get("faculty_type", "custom")).strip(), salary, str(data.get("joining_date", "")).strip(), str(data.get("relieving_date", "")).strip(), str(data.get("email", "")).strip(), str(data.get("mobile", "")).strip(), json.dumps(weekdays)),
            )
        return jsonify({"ok": True, "id": cursor.lastrowid}), 201
    except (ValueError, TypeError):
        return error("Monthly salary and weekdays must be valid")
    except sqlite3.IntegrityError:
        return error("Employee code already exists or department is invalid", 409)


@app.patch("/api/faculties/<int:faculty_id>")
def update_faculty(faculty_id):
    data = request.get_json(silent=True) or {}
    required = ["employee_code", "name", "department_id", "monthly_salary"]
    if any(str(data.get(field, "")).strip() == "" for field in required):
        return error("Employee code, name, department, and monthly salary are required")
    try:
        weekdays = sorted({int(day) for day in data.get("weekdays", []) if 0 <= int(day) <= 6})
        salary = float(data["monthly_salary"])
        if salary < 0:
            raise ValueError
        with db() as connection:
            cursor = connection.execute(
                "UPDATE faculties SET employee_code=?, name=?, department_id=?, designation=?, faculty_type=?, monthly_salary=?, joining_date=?, relieving_date=?, email=?, mobile=?, weekdays=? WHERE id=? AND status='active'",
                (str(data["employee_code"]).strip().upper(), str(data["name"]).strip(), int(data["department_id"]), str(data.get("designation", "")).strip(), str(data.get("faculty_type", "custom")).strip(), salary, str(data.get("joining_date", "")).strip(), str(data.get("relieving_date", "")).strip(), str(data.get("email", "")).strip(), str(data.get("mobile", "")).strip(), json.dumps(weekdays), faculty_id),
            )
        if cursor.rowcount == 0:
            return error("Faculty member not found", 404)
        sync_existing_calendar_weekdays(faculty_id, weekdays)
        return jsonify({"ok": True})
    except (ValueError, TypeError):
        return error("Monthly salary and weekdays must be valid")
    except sqlite3.IntegrityError:
        return error("Employee code already exists or department is invalid", 409)


@app.delete("/api/faculties/<int:faculty_id>")
def delete_faculty(faculty_id):
    with db() as connection:
        cursor = connection.execute("UPDATE faculties SET status='inactive' WHERE id=? AND status='active'", (faculty_id,))
    if cursor.rowcount == 0:
        return error("Faculty member not found", 404)
    return jsonify({"ok": True})


@app.get("/api/attendance")
def get_attendance():
    faculty_id = request_int(request.args.get("faculty_id"), 0)
    month = request_int(request.args.get("month"), date.today().month)
    year = request_int(request.args.get("year"), date.today().year)
    if faculty_id <= 0 or month not in range(1, 13) or year < 2000:
        return error("Faculty, month, and year are required")
    faculty, rows = ensure_calendar(faculty_id, month, year)
    if not faculty:
        return error("Faculty not found", 404)
    return jsonify({"faculty": faculty_dict(faculty), "attendance": rows, "totals": attendance_totals(rows)})


@app.put("/api/attendance/<int:attendance_id>")
def update_attendance(attendance_id):
    data = request.get_json(silent=True) or {}
    is_working = int(bool(data.get("is_working")))
    requested_status = str(data.get("attendance_status", "")).strip().lower()
    status = requested_status if is_working or requested_status == "special" else ""
    leave_type = str(data.get("leave_type", "")).strip().lower() if status == "leave" else ""
    if status not in {"", "present", "absent", "leave", "special"} or (status == "special" and is_working) or leave_type not in {"", "paid", "unpaid"}:
        return error("Attendance or leave type is invalid")
    with db() as connection:
        cursor = connection.execute("UPDATE monthly_attendance SET is_working=?, attendance_status=?, leave_type=?, remarks=? WHERE id=?", (is_working, status, leave_type, str(data.get("remarks", "")).strip(), attendance_id))
    if cursor.rowcount == 0:
        return error("Attendance date not found", 404)
    return jsonify({"ok": True})


@app.post("/api/salary/calculate")
def calculate():
    data = request.get_json(silent=True) or {}
    faculty_id = request_int(data.get("faculty_id"), 0)
    month, year = request_int(data.get("month"), 0), request_int(data.get("year"), 0)
    faculty, rows = ensure_calendar(faculty_id, month, year) if month in range(1, 13) else (None, [])
    if not faculty:
        return error("Faculty and valid month are required")
    totals = attendance_totals(rows)
    try:
        result = calculate_salary(
            monthly_salary=faculty["monthly_salary"],
            adjustment=data.get("adjustment", 0),
            working_days=totals["working_days"],
            present_days=totals["regular_present_days"],
            paid_leave=totals["paid_leave"],
            unpaid_leave=totals["unpaid_leave"],
            special_present_days=totals["special_present_days"],
        )
    except ValueError as exc:
        return error(str(exc))
    snapshot = {key: str(value) for key, value in result.items()}
    with db() as connection:
        connection.execute("INSERT INTO salary_records (faculty_id, month, year, snapshot) VALUES (?, ?, ?, ?) ON CONFLICT(faculty_id, month, year) DO UPDATE SET snapshot=excluded.snapshot, status='draft'", (faculty_id, month, year, json.dumps(snapshot)))
    return jsonify({"ok": True, "salary": snapshot, "totals": totals})


@app.post("/api/salary/<int:faculty_id>/finalize")
def finalize_salary(faculty_id):
    data = request.get_json(silent=True) or {}
    month, year = request_int(data.get("month"), 0), request_int(data.get("year"), 0)
    with db() as connection:
        cursor = connection.execute("UPDATE salary_records SET status='finalized' WHERE faculty_id=? AND month=? AND year=?", (faculty_id, month, year))
        if cursor.rowcount == 0:
            return error("Calculate salary before finalizing", 409)
        connection.execute("INSERT INTO audit_logs (action, module, record_id) VALUES ('Salary finalized', 'salary', ?)", (str(faculty_id),))
    return jsonify({"ok": True, "status": "finalized"})


init_db()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5001, debug=True)
